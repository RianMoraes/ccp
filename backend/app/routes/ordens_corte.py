import hashlib
import json
import re
import unicodedata
import warnings
from collections import Counter
from datetime import date, datetime
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlmodel import Session, SQLModel, select

from app.database import get_session
from app.models import (
    Componente, Equipamento, HistoricoOrdemCorte, OrdemCorte,
    ReferenciaRetrabalhoCorte, StatusOrdemCorteEnum, TipoOrdemCorteEnum,
    Usuario,
)
from app.routes.auth import exigir_operacao, obter_usuario_atual


router = APIRouter(prefix="/api/ordens-corte", tags=["Ordens de Corte"])


class OrdemCorteCreate(SQLModel):
    numero: str
    tipo: TipoOrdemCorteEnum
    material: str | None = None
    material_original: str | None = None
    material_grupo: str | None = None
    prioridade: str | None = None
    data_insercao: str | None = None
    tipo_registro: str | None = None
    vinculo_rtb: str | None = None
    liga: str | None = None
    espessura: str | None = None
    dimensao_x: str | None = None
    dimensao_y: str | None = None
    quantidade: int | None = None
    material_completo: str | None = None
    op: str | None = None
    rv: str | None = None
    equipamento_id: str | None = None
    componente_id: str | None = None
    equipamento_texto: str | None = None
    componente_texto: str | None = None
    maquina: str | None = None
    corte_planilha: str | None = None
    status_planilha: str | None = None
    observacoes: str | None = None
    status: StatusOrdemCorteEnum = StatusOrdemCorteEnum.PROGRAMADO
    data_corte: datetime | None = None
    motivo_bloqueio: str | None = None
    eh_retrabalho: bool = False
    pagina_inicial: int = 1
    quantidade_paginas: int = 1
    gerar_paginas: bool = False


class OrdemCorteUpdate(OrdemCorteCreate):
    ordem_base: str | None = None
    pagina: int | None = None
    total_paginas: int | None = None


class AlteracaoStatusCorte(SQLModel):
    status: StatusOrdemCorteEnum
    motivo: str | None = None


class RetrabalhoCreate(SQLModel):
    numero: str
    ordens_originais_ids: list[str]
    motivo: str
    responsavel: str | None = None


class StatusEmMassa(SQLModel):
    ids: list[str]
    status: StatusOrdemCorteEnum
    motivo: str | None = None


class ExclusaoEmMassa(SQLModel):
    ids: list[str] = []
    apagar_todas: bool = False
    confirmacao: str | None = None


def _texto(valor):
    if valor is None:
        return None
    valor = str(valor).strip()
    return None if not valor or valor == "-" else valor


def _chave_texto(valor):
    valor = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", " ", valor).strip()


def _normalizar_material(valor):
    original = _texto(valor)
    chave = _chave_texto(original)
    aliases = {
        "ACO RED": "Aço redondo", "ACO REDONDO": "Aço redondo", "ACO REDONDO TREFILADO": "Aço redondo",
        "ACO REDONDO LAMINADO": "Aço redondo", "ACO REDONDO INOX": "Aço redondo",
        "ACO QUA": "Aço quadrado", "ACO QUAD": "Aço quadrado", "F QUAD": "Aço quadrado", "F QUA": "Aço quadrado",
        "ACO SEXT": "Aço sextavado", "ACO SEXTAVADO": "Aço sextavado",
        "F CH": "Aço chato", "F CHATO": "Aço chato", "ACO CHATO": "Aço chato",
        "TUBO": "Tubo", "TUBO QUADRADO": "Tubo", "TUBO INDUSTRIAL": "Tubo", "TUBO MECANICO": "Tubo",
        "TUBO DIN2440": "Tubo", "TUBO DIN 2440": "Tubo", "TUBO SCH 20": "Tubo", "TUBO SCH20": "Tubo",
        "TUBO SCH 40": "Tubo", "TUBO SCH40": "Tubo", "TUBO SCH 80": "Tubo", "TUBO SCH80": "Tubo",
        "TUBO SCH 80S": "Tubo", "TUBO SCH80S": "Tubo", "TUBO SCH10S": "Tubo",
        "CANTONEIRA": "Cantoneira", "CHAPA": "Chaparia", "CHAPARIA": "Chaparia",
        "UHMW": "UHMW", "POLICARBONATO": "Policarbonato", "BORRACHA": "Borracha",
        "BRONZE": "Bronze", "VIGA": "Viga", "VIGA H": "Viga", "VIGA I": "Viga", "VIGA U": "Viga", "VIGA W": "Viga",
        "FERRO FUNDIDO": "Ferro fundido", "F F": "Ferro fundido",
        "TELA": "Tela", "VERGALHAO": "Vergalhão", "METALON": "Metalon", "PERFILADOS": "Perfilados",
        "SAC350": "SAC350", "AR400": "AR400", "AR450": "AR450", "HARDOX": "Hardox",
        "NYLON": "Nylon", "TECHNYL": "Nylon", "DISCO DE CHAPA": "Disco de chapa",
    }
    if chave == "VIGA CANT":
        return "Barras e perfis", "Viga/Cantoneira", original, False
    material = aliases.get(chave, original)
    grupos = {
        "Aço redondo": "Barras e perfis", "Aço quadrado": "Barras e perfis", "Aço sextavado": "Barras e perfis",
        "Aço chato": "Barras e perfis", "Cantoneira": "Barras e perfis", "Viga": "Barras e perfis",
        "Vergalhão": "Barras e perfis", "Metalon": "Barras e perfis", "Perfilados": "Barras e perfis",
        "Tubo": "Tubos", "Chaparia": "Chapas", "Tela": "Chapas", "Disco de chapa": "Chapas",
        "SAC350": "Chapas especiais", "AR400": "Chapas especiais", "AR450": "Chapas especiais", "Hardox": "Chapas especiais",
        "UHMW": "Polímeros e borrachas", "Policarbonato": "Polímeros e borrachas", "Borracha": "Polímeros e borrachas", "Nylon": "Polímeros e borrachas",
        "Bronze": "Metais especiais", "Ferro fundido": "Metais especiais",
    }
    return grupos.get(material, "Outros"), material, original, material not in grupos


def _inferir_tipo(material, maquina):
    maquina_chave = _chave_texto(maquina)
    if maquina_chave in {"SERRA", "GEKA", "PERFIL"}:
        return TipoOrdemCorteEnum.CORTE_FRIO
    if maquina_chave in {"DARDI", "WELLE", "ROUTER", "LASER", "OXICORTE", "GUILHOTINA", "TNL", "WE"}:
        return TipoOrdemCorteEnum.CNC
    if material in {"Chaparia", "Tela", "Disco de chapa", "SAC350", "AR400", "AR450", "Hardox", "UHMW", "Policarbonato", "Borracha", "Nylon"}:
        return TipoOrdemCorteEnum.CNC
    return TipoOrdemCorteEnum.CORTE_FRIO


def _separar_pagina_cnc(numero):
    correspondencia = re.fullmatch(r"(.+)-(\d{3})", str(numero or "").strip())
    if not correspondencia:
        return str(numero or "").strip(), None
    return correspondencia.group(1), int(correspondencia.group(2))


def _ler_excel(conteudo):
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
            workbook = openpyxl.load_workbook(BytesIO(conteudo), read_only=False, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Não foi possível abrir o arquivo Excel.") from exc
    if "ORDENS DE CORTE" not in workbook.sheetnames:
        raise HTTPException(status_code=422, detail="A aba 'ORDENS DE CORTE' não foi encontrada.")
    ws = workbook["ORDENS DE CORTE"]
    linha_cabecalho = None
    cabecalhos = None
    for numero_linha, valores in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True), 1):
        chaves = [_chave_texto(valor) for valor in valores]
        if "ORDEM DE CORTE" in chaves or "ORDEM" in chaves:
            linha_cabecalho = numero_linha
            cabecalhos = {chave: indice for indice, chave in enumerate(chaves) if chave}
            break
    if not linha_cabecalho:
        raise HTTPException(status_code=422, detail="Não foi possível localizar o cabeçalho das ordens de corte.")

    def obter(valores, *nomes):
        for nome in nomes:
            indice = cabecalhos.get(_chave_texto(nome))
            if indice is not None and indice < len(valores):
                return valores[indice]
        return None

    def texto_data(valor):
        if isinstance(valor, (datetime, date)):
            return valor.isoformat()
        return _texto(valor)

    def inteiro(valor):
        if valor is None or valor == "":
            return None
        try:
            return int(float(str(valor).replace(",", ".")))
        except (TypeError, ValueError):
            return None

    ordens = []
    desconhecidos = Counter()
    for linha, valores in enumerate(ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True), linha_cabecalho + 1):
        numero = _texto(obter(valores, "ORDEM DE CORTE", "ORDEM"))
        if not numero:
            continue
        grupo, material, material_original, desconhecido = _normalizar_material(obter(valores, "MATERIAL"))
        if desconhecido and material_original:
            desconhecidos[material_original] += 1
        corte_planilha = _chave_texto(obter(valores, "CORTE (S/N)", "CORTE")) or None
        status_planilha = _texto(obter(valores, "STATUS DE CORTE", "STATUS"))
        status_chave = _chave_texto(status_planilha)
        if corte_planilha == "S" or status_chave == "CORTADO":
            status_normalizado = StatusOrdemCorteEnum.CONCLUIDO
        elif corte_planilha == "R" or status_chave == "RETRABALHADO":
            status_normalizado = StatusOrdemCorteEnum.RETRABALHADO
        elif corte_planilha == "C" or status_chave == "CANCELADO":
            status_normalizado = StatusOrdemCorteEnum.CANCELADO
        elif status_chave == "SEM MATERIAL":
            status_normalizado = StatusOrdemCorteEnum.BLOQUEADO
        else:
            status_normalizado = StatusOrdemCorteEnum.PROGRAMADO
        data_corte_valor = obter(valores, "DATA DE CORTE", "DATA CORTE")
        data_corte = texto_data(data_corte_valor) if isinstance(data_corte_valor, (datetime, date)) else None
        maquina = _texto(obter(valores, "MÁQUINA", "MAQUINA"))
        tipo_registro = _texto(obter(valores, "TIPO")) or "ORIGINAL"
        liga = _texto(obter(valores, "LIGA", "DUREZA"))
        espessura = _texto(obter(valores, "ESP", "ESPESSURA"))
        dimensao_x = _texto(obter(valores, "DIM X", "DIMENSÃO X"))
        dimensao_y = _texto(obter(valores, "DIM Y", "DIMENSÃO Y"))
        quantidade = inteiro(obter(valores, "QNT", "QUANTIDADE"))
        material_completo = _texto(obter(valores, "MATERIAL COMPLETO PARA TABELA"))
        if not material_completo:
            partes = [parte for parte in [material_original, liga, espessura, dimensao_x, dimensao_y] if parte]
            material_completo = " · ".join(partes) or None
        bruto = {
            "numero": numero,
            "prioridade": _texto(obter(valores, "PRIORIDADE")),
            "data_insercao": texto_data(obter(valores, "DATA DE INSERÇÃO DE ORDEM NA PLANILHA", "DATA DE INSERCAO DE ORDEM NA PLANILHA")),
            "tipo_registro": tipo_registro,
            "vinculo_rtb": _texto(obter(valores, "VÍNCULO RTB", "VINCULO RTB")),
            "material_original": material_original,
            "liga": liga,
            "espessura": espessura,
            "dimensao_x": dimensao_x,
            "dimensao_y": dimensao_y,
            "quantidade": quantidade,
            "material_completo": material_completo,
            "rv": _texto(obter(valores, "RV TSG", "RV")),
            "op": _texto(obter(valores, "OP TSG", "OP")),
            "equipamento_texto": _texto(obter(valores, "EQUIPAMENTO")),
            "componente_texto": _texto(obter(valores, "COMPONENTE")),
            "corte_planilha": corte_planilha,
            "status_planilha": status_planilha,
            "data_corte": data_corte,
            "maquina": maquina,
            "observacoes": _texto(obter(valores, "OBS", "OBSERVAÇÕES", "OBSERVACOES")),
            "linha": linha,
        }
        identidade = {chave: bruto.get(chave) for chave in [
            "numero", "tipo_registro", "material_original", "liga", "espessura",
            "dimensao_x", "dimensao_y", "quantidade", "op", "componente_texto",
            "data_insercao",
        ]}
        chave = "excel-ordem:" + hashlib.sha256(json.dumps(identidade, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        ordens.append({
            **bruto, "material_grupo": grupo, "material": material,
            "tipo": _inferir_tipo(material, maquina or status_planilha),
            "status": status_normalizado,
            "eh_retrabalho": _chave_texto(tipo_registro) == "RETRABALHO",
            "motivo_bloqueio": "Sem material (sincronizado da planilha)" if status_normalizado == StatusOrdemCorteEnum.BLOQUEADO else None,
            "chave_importacao": chave,
        })

    grupos_paginas = {}
    for item in ordens:
        if item["tipo"] != TipoOrdemCorteEnum.CNC:
            item.update({"ordem_base": None, "pagina": None, "total_paginas": None})
            continue
        base, pagina = _separar_pagina_cnc(item["numero"])
        item["ordem_base"] = base
        item["pagina"] = pagina
        if pagina is not None:
            grupos_paginas.setdefault(base, []).append(pagina)
    for item in ordens:
        paginas = grupos_paginas.get(item.get("ordem_base"), [])
        item["total_paginas"] = max(paginas) if item.get("pagina") is not None and paginas else item.get("pagina")
    ordens_unicas = {}
    for item in ordens:
        ordens_unicas[item["chave_importacao"]] = item
    return list(ordens_unicas.values()), desconhecidos


def _registrar(session, ordem_id, usuario, evento, descricao):
    session.add(HistoricoOrdemCorte(ordem_corte_id=ordem_id, usuario=usuario, evento=evento, descricao=descricao))


def _aplicar_status(ordem, novo_status, motivo=None):
    if novo_status == ordem.status:
        if novo_status == StatusOrdemCorteEnum.BLOQUEADO and (motivo or "").strip():
            ordem.motivo_bloqueio = motivo.strip()
            ordem.atualizado_em = datetime.utcnow()
        return
    if novo_status == StatusOrdemCorteEnum.BLOQUEADO:
        if not (motivo or "").strip():
            raise HTTPException(status_code=400, detail="Informe o motivo do bloqueio.")
        ordem.status_anterior = ordem.status
        ordem.motivo_bloqueio = motivo.strip()
    elif ordem.status == StatusOrdemCorteEnum.BLOQUEADO:
        ordem.motivo_bloqueio = None
        ordem.status_anterior = None
    if novo_status == StatusOrdemCorteEnum.CONCLUIDO and not ordem.data_corte:
        ordem.data_corte = datetime.utcnow()
    ordem.status = novo_status
    ordem.atualizado_em = datetime.utcnow()


def _serializar_ordens(session, ordens):
    referencias = session.exec(select(ReferenciaRetrabalhoCorte)).all()
    refs_por_rtb = {}
    rtbs_por_original = {}
    for ref in referencias:
        refs_por_rtb.setdefault(ref.retrabalho_ordem_id, []).append(ref)
        if ref.ordem_original_id:
            rtbs_por_original.setdefault(ref.ordem_original_id, []).append(ref.retrabalho_ordem_id)
    por_id = {ordem.id: ordem for ordem in ordens}
    equipamentos = {item.id: item for item in session.exec(select(Equipamento)).all()}
    componentes = {item.id: item for item in session.exec(select(Componente)).all()}
    paginas_por_base = {}
    for ordem in ordens:
        if ordem.tipo != TipoOrdemCorteEnum.CNC:
            continue
        base = ordem.ordem_base
        pagina = ordem.pagina
        if pagina is None:
            base, pagina = _separar_pagina_cnc(ordem.numero)
        if base and pagina is not None:
            paginas_por_base.setdefault(base, []).append(pagina)
    resultado = []
    for ordem in ordens:
        dados = ordem.model_dump()
        if ordem.tipo == TipoOrdemCorteEnum.CNC and dados.get("pagina") is None:
            base, pagina = _separar_pagina_cnc(ordem.numero)
            if pagina is not None:
                dados["ordem_base"] = base
                dados["pagina"] = pagina
                dados["total_paginas"] = max(paginas_por_base.get(base, [pagina]))
        dados["equipamento_nome"] = equipamentos[ordem.equipamento_id].nome if ordem.equipamento_id in equipamentos else ordem.equipamento_texto
        dados["componente_nome"] = componentes[ordem.componente_id].nome if ordem.componente_id in componentes else ordem.componente_texto
        dados["referencias_originais"] = [
            {"ordem_id": ref.ordem_original_id, "texto": ref.trabalho_original_texto}
            for ref in refs_por_rtb.get(ordem.id, [])
        ]
        dados["retrabalhos"] = [
            {"id": rtb_id, "numero": por_id[rtb_id].numero}
            for rtb_id in rtbs_por_original.get(ordem.id, []) if rtb_id in por_id
        ]
        resultado.append(dados)
    return resultado


@router.get("")
def listar_ordens(session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    ordens = session.exec(select(OrdemCorte).order_by(OrdemCorte.criado_em.desc())).all()
    return _serializar_ordens(session, ordens)


@router.post("", status_code=status.HTTP_201_CREATED, dependencies=[Depends(exigir_operacao)])
def criar_ordem(dados: OrdemCorteCreate, session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    numero = dados.numero.strip()
    if not numero:
        raise HTTPException(status_code=400, detail="Informe o número da ordem.")
    quantidade_paginas = max(int(dados.quantidade_paginas or 1), 1)
    pagina_inicial = max(int(dados.pagina_inicial or 1), 1)
    if quantidade_paginas > 1000:
        raise HTTPException(status_code=400, detail="O limite por geração é de 1.000 páginas.")
    base, pagina_numero = _separar_pagina_cnc(numero)
    if dados.tipo == TipoOrdemCorteEnum.CNC:
        if dados.gerar_paginas:
            base = numero.rstrip("-")
            numeros = [f"{base}-{pagina:03d}" for pagina in range(pagina_inicial, pagina_inicial + quantidade_paginas)]
        else:
            numeros = [numero]
            pagina_inicial = pagina_numero
            quantidade_paginas = 1
    else:
        numeros = [numero]
        base = None
        pagina_inicial = None
        quantidade_paginas = 1
    existentes = session.exec(select(OrdemCorte).where(OrdemCorte.numero.in_(numeros))).all()
    if existentes:
        raise HTTPException(status_code=409, detail=f"Já existem ordens com estes números: {', '.join(item.numero for item in existentes[:10])}.")
    grupo, material, original, _ = _normalizar_material(dados.material)
    criadas = []
    for indice, numero_gerado in enumerate(numeros):
        pagina = pagina_inicial + indice if dados.tipo == TipoOrdemCorteEnum.CNC and pagina_inicial is not None else None
        ordem = OrdemCorte(
            **dados.model_dump(exclude={"numero", "material", "material_original", "material_grupo", "pagina_inicial", "quantidade_paginas", "gerar_paginas"}),
            numero=numero_gerado, ordem_base=base, pagina=pagina,
            total_paginas=(pagina_inicial + quantidade_paginas - 1) if dados.tipo == TipoOrdemCorteEnum.CNC and pagina_inicial is not None else None,
            material=material, material_grupo=dados.material_grupo or grupo, material_original=dados.material_original or original, criado_por=usuario_atual.nome,
        )
        session.add(ordem)
        session.flush()
        _registrar(session, ordem.id, usuario_atual.nome, "Ordem criada", f"Ordem {numero_gerado} criada como programada.")
        criadas.append(ordem)
    session.commit()
    if len(criadas) == 1:
        session.refresh(criadas[0])
        return criadas[0]
    return {"message": f"{len(criadas)} páginas CNC criadas.", "quantidade": len(criadas), "primeira": criadas[0].numero, "ultima": criadas[-1].numero}


@router.put("/{ordem_id}", dependencies=[Depends(exigir_operacao)])
def editar_ordem(ordem_id: str, dados: OrdemCorteUpdate, session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    ordem = session.get(OrdemCorte, ordem_id)
    if not ordem:
        raise HTTPException(status_code=404, detail="Ordem de corte não encontrada.")
    numero = dados.numero.strip()
    if not numero:
        raise HTTPException(status_code=400, detail="Informe o número da ordem.")
    grupo, material, original, _ = _normalizar_material(dados.material)
    status_original = ordem.status
    for campo in [
        "tipo", "prioridade", "data_insercao", "tipo_registro", "vinculo_rtb", "liga", "espessura",
        "dimensao_x", "dimensao_y", "quantidade", "material_completo", "op", "rv", "equipamento_texto",
        "componente_texto", "maquina", "corte_planilha", "status_planilha", "observacoes", "data_corte",
        "eh_retrabalho",
    ]:
        setattr(ordem, campo, getattr(dados, campo))
    equipamento = session.exec(select(Equipamento).where(Equipamento.op == dados.op)).first() if dados.op else None
    if not equipamento and dados.equipamento_texto:
        equipamento = session.exec(select(Equipamento).where(Equipamento.nome == dados.equipamento_texto)).first()
    ordem.equipamento_id = equipamento.id if equipamento else None
    componente = session.exec(select(Componente).where(
        Componente.equipamento_id == equipamento.id,
        Componente.nome == dados.componente_texto,
    )).first() if equipamento and dados.componente_texto else None
    ordem.componente_id = componente.id if componente else None
    ordem.numero = numero
    ordem.material = material
    ordem.material_original = dados.material_original or original
    ordem.material_grupo = dados.material_grupo or grupo
    ordem.ordem_base = (dados.ordem_base or "").strip() or None
    ordem.pagina = dados.pagina
    ordem.total_paginas = dados.total_paginas
    if dados.status != status_original:
        _aplicar_status(ordem, dados.status, dados.motivo_bloqueio)
    elif ordem.status == StatusOrdemCorteEnum.BLOQUEADO:
        ordem.motivo_bloqueio = (dados.motivo_bloqueio or "").strip() or ordem.motivo_bloqueio
    ordem.atualizado_em = datetime.utcnow()
    session.add(ordem)
    _registrar(session, ordem.id, usuario_atual.nome, "Ordem editada", "Dados da ordem atualizados manualmente.")
    session.commit()
    return {"message": "Ordem atualizada."}


@router.patch("/{ordem_id}/status", dependencies=[Depends(exigir_operacao)])
def alterar_status(ordem_id: str, dados: AlteracaoStatusCorte, session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    ordem = session.get(OrdemCorte, ordem_id)
    if not ordem:
        raise HTTPException(status_code=404, detail="Ordem de corte não encontrada.")
    _aplicar_status(ordem, dados.status, dados.motivo)
    session.add(ordem)
    _registrar(session, ordem.id, usuario_atual.nome, "Status alterado", f"Ordem alterada para {dados.status.value}." + (f" Motivo: {dados.motivo}" if dados.motivo else ""))
    session.commit()
    return {"message": "Status atualizado."}


@router.patch("/status-em-massa", dependencies=[Depends(exigir_operacao)])
def alterar_status_em_massa(dados: StatusEmMassa, session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    ordens = [session.get(OrdemCorte, item) for item in set(dados.ids)]
    if not ordens or any(item is None for item in ordens):
        raise HTTPException(status_code=400, detail="Selecione ordens válidas.")
    for ordem in ordens:
        _aplicar_status(ordem, dados.status, dados.motivo)
    for ordem in ordens:
        session.add(ordem)
        _registrar(session, ordem.id, usuario_atual.nome, "Status alterado em massa", f"Ordem alterada para {dados.status.value}.")
    session.commit()
    return {"message": f"{len(ordens)} ordem(ns) atualizada(s).", "quantidade": len(ordens)}


@router.post("/excluir-em-massa", dependencies=[Depends(exigir_operacao)])
def excluir_em_massa(dados: ExclusaoEmMassa, session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    if dados.apagar_todas:
        if dados.confirmacao != "APAGAR TODAS":
            raise HTTPException(status_code=400, detail="Digite APAGAR TODAS para confirmar.")
        ordens = session.exec(select(OrdemCorte)).all()
    else:
        ids = set(dados.ids)
        if not ids:
            raise HTTPException(status_code=400, detail="Selecione ao menos uma ordem.")
        ordens = session.exec(select(OrdemCorte).where(OrdemCorte.id.in_(ids))).all()
    ids_ordens = {item.id for item in ordens}
    if not ids_ordens:
        return {"message": "Nenhuma ordem para excluir.", "quantidade": 0}
    todas_referencias = session.exec(select(ReferenciaRetrabalhoCorte)).all()
    referencias = [item for item in todas_referencias if
        item.retrabalho_ordem_id in ids_ordens or item.ordem_original_id in ids_ordens]
    originais_afetadas = {item.ordem_original_id for item in referencias if item.retrabalho_ordem_id in ids_ordens and item.ordem_original_id and item.ordem_original_id not in ids_ordens}
    for original_id in originais_afetadas:
        possui_outro_rtb = any(
            item.ordem_original_id == original_id and item.retrabalho_ordem_id not in ids_ordens
            for item in todas_referencias
        )
        original = session.get(OrdemCorte, original_id)
        if original and not possui_outro_rtb and original.status == StatusOrdemCorteEnum.RETRABALHADO:
            original.status = original.status_anterior or StatusOrdemCorteEnum.PROGRAMADO
            original.status_anterior = None
            original.atualizado_em = datetime.utcnow()
            session.add(original)
    historicos = session.exec(select(HistoricoOrdemCorte).where(HistoricoOrdemCorte.ordem_corte_id.in_(ids_ordens))).all()
    for item in referencias + historicos:
        session.delete(item)
    for ordem in ordens:
        session.delete(ordem)
    session.commit()
    return {"message": f"{len(ordens)} ordem(ns) excluída(s).", "quantidade": len(ordens)}


@router.post("/retrabalhos", status_code=status.HTTP_201_CREATED, dependencies=[Depends(exigir_operacao)])
def criar_retrabalho(dados: RetrabalhoCreate, session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    numero = dados.numero.strip().upper()
    if not numero:
        raise HTTPException(status_code=400, detail="Informe o número ou identificação do retrabalho.")
    if session.exec(select(OrdemCorte).where(OrdemCorte.numero == numero)).first():
        raise HTTPException(status_code=409, detail="Este número de retrabalho já existe.")
    originais = [session.get(OrdemCorte, item) for item in dados.ordens_originais_ids]
    if not originais or any(item is None for item in originais):
        raise HTTPException(status_code=400, detail="Selecione ao menos uma ordem original válida.")
    rtb = OrdemCorte(numero=numero, tipo=originais[0].tipo, material_grupo=originais[0].material_grupo, material=originais[0].material, status=StatusOrdemCorteEnum.PROGRAMADO, eh_retrabalho=True, responsavel_retrabalho=dados.responsavel, observacoes=dados.motivo, origem="manual", criado_por=usuario_atual.nome)
    session.add(rtb)
    session.flush()
    for original in originais:
        if original.status != StatusOrdemCorteEnum.RETRABALHADO:
            original.status_anterior = original.status
        original.status = StatusOrdemCorteEnum.RETRABALHADO
        original.atualizado_em = datetime.utcnow()
        session.add(original)
        session.add(ReferenciaRetrabalhoCorte(retrabalho_ordem_id=rtb.id, ordem_original_id=original.id, trabalho_original_texto=original.numero))
        _registrar(session, original.id, usuario_atual.nome, "Retrabalho aberto", f"Vinculada ao retrabalho {numero}: {dados.motivo}")
    _registrar(session, rtb.id, usuario_atual.nome, "Retrabalho criado", f"Criado a partir de {len(originais)} ordem(ns).")
    session.commit()
    return {"id": rtb.id, "numero": numero}


@router.post("/importar/analisar", dependencies=[Depends(exigir_operacao)])
async def analisar_excel(arquivo: UploadFile = File(...), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    conteudo = await arquivo.read()
    ordens, desconhecidos = _ler_excel(conteudo)
    numeros = Counter(item["numero"] for item in ordens)
    ops = {
        parte.strip()
        for item in ordens
        for parte in re.split(r"\s*;\s*", item.get("op") or "")
        if parte.strip()
    }
    return {
        "arquivo": arquivo.filename, "ordens": len(ordens), "concluidas": sum(item["status"] == StatusOrdemCorteEnum.CONCLUIDO for item in ordens),
        "programadas": sum(item["status"] == StatusOrdemCorteEnum.PROGRAMADO for item in ordens),
        "bloqueadas": sum(item["status"] == StatusOrdemCorteEnum.BLOQUEADO for item in ordens),
        "pendentes": sum(item["status"] in {StatusOrdemCorteEnum.PROGRAMADO, StatusOrdemCorteEnum.EM_CORTE, StatusOrdemCorteEnum.BLOQUEADO} for item in ordens),
        "ops": len(ops), "retrabalhos": sum(item["eh_retrabalho"] for item in ordens),
        "duplicidades": [{"numero": numero, "quantidade": qtd} for numero, qtd in numeros.items() if qtd > 1],
        "materiais_desconhecidos": [{"material": nome, "quantidade": qtd} for nome, qtd in desconhecidos.most_common()],
        "amostra": ordens[:100],
    }


@router.post("/importar/confirmar", dependencies=[Depends(exigir_operacao)])
async def confirmar_excel(arquivo: UploadFile = File(...), session: Session = Depends(get_session), usuario_atual: Usuario = Depends(obter_usuario_atual)):
    conteudo = await arquivo.read()
    ordens, _ = _ler_excel(conteudo)
    existentes = session.exec(select(OrdemCorte)).all()
    por_chave = {item.chave_importacao: item for item in existentes if item.chave_importacao}
    existentes_por_numero = {}
    for item in existentes:
        existentes_por_numero.setdefault(item.numero.casefold(), []).append(item)
    numeros_planilha = Counter(item["numero"].casefold() for item in ordens)
    equipamentos = session.exec(select(Equipamento).where(Equipamento.ativo == True)).all()
    equipamento_por_op = {str(item.op).strip().casefold(): item for item in equipamentos if item.op}
    componentes = session.exec(select(Componente).where(Componente.ativo == True)).all()
    componentes_por_equip_nome = {(item.equipamento_id, _chave_texto(item.nome)): item for item in componentes}
    criadas = 0
    atualizadas = 0
    sem_alteracao = 0
    ids_sincronizados = set()
    agora = datetime.utcnow()
    for item in ordens:
        numero_chave = item["numero"].casefold()
        ordem = por_chave.get(item["chave_importacao"])
        if ordem and ordem.id in ids_sincronizados:
            ordem = None
        candidatos = [
            candidato for candidato in existentes_por_numero.get(numero_chave, [])
            if candidato.id not in ids_sincronizados
        ]
        if not ordem and candidatos:
            compativeis = [
                candidato for candidato in candidatos
                if _chave_texto(candidato.material_original or candidato.material) == _chave_texto(item.get("material_original") or item.get("material"))
                and _chave_texto(candidato.componente_texto) == _chave_texto(item.get("componente_texto"))
            ]
            if len(compativeis) == 1:
                ordem = compativeis[0]
        if not ordem and numeros_planilha[numero_chave] == 1 and len(candidatos) == 1:
            ordem = candidatos[0]
        ops_item = [parte.strip().casefold() for parte in re.split(r"\s*;\s*", item.get("op") or "") if parte.strip()]
        equipamento = next((equipamento_por_op.get(op) for op in ops_item if equipamento_por_op.get(op)), None)
        componente = componentes_por_equip_nome.get((equipamento.id, _chave_texto(item.get("componente_texto")))) if equipamento else None
        valores = {
            "numero": item["numero"], "tipo": item["tipo"], "prioridade": item.get("prioridade"),
            "data_insercao": item.get("data_insercao"), "tipo_registro": item.get("tipo_registro"),
            "vinculo_rtb": item.get("vinculo_rtb"), "material_grupo": item["material_grupo"],
            "material": item["material"], "material_original": item["material_original"], "liga": item.get("liga"),
            "espessura": item.get("espessura"), "dimensao_x": item.get("dimensao_x"), "dimensao_y": item.get("dimensao_y"),
            "quantidade": item.get("quantidade"), "material_completo": item.get("material_completo"),
            "ordem_base": item.get("ordem_base"), "pagina": item.get("pagina"), "total_paginas": item.get("total_paginas"),
            "rv": item["rv"], "op": item["op"], "equipamento_id": equipamento.id if equipamento else None,
            "componente_id": componente.id if componente else None, "equipamento_texto": item["equipamento_texto"],
            "componente_texto": item["componente_texto"], "status": item["status"],
            "data_corte": datetime.fromisoformat(item["data_corte"]) if item["data_corte"] else None,
            "maquina": item["maquina"], "corte_planilha": item.get("corte_planilha"),
            "status_planilha": item.get("status_planilha"), "observacoes": item["observacoes"],
            "motivo_bloqueio": item.get("motivo_bloqueio"), "eh_retrabalho": item.get("eh_retrabalho", False),
            "origem": "excel", "linha_origem": item["linha"], "chave_importacao": item["chave_importacao"],
            "presente_planilha": True,
        }
        if ordem:
            alterou = any(getattr(ordem, campo) != valor for campo, valor in valores.items())
            for campo, valor in valores.items():
                setattr(ordem, campo, valor)
            ordem.atualizado_em = agora
            ordem.sincronizado_em = agora
            session.add(ordem)
            if alterou:
                atualizadas += 1
                _registrar(session, ordem.id, usuario_atual.nome, "Ordem sincronizada", f"Dados atualizados pela planilha {arquivo.filename}.")
            else:
                sem_alteracao += 1
        else:
            ordem = OrdemCorte(**valores, criado_por=usuario_atual.nome, sincronizado_em=agora)
            session.add(ordem)
            session.flush()
            criadas += 1
            _registrar(session, ordem.id, usuario_atual.nome, "Ordem importada", f"Criada pela planilha {arquivo.filename}.")
        por_chave[item["chave_importacao"]] = ordem
        ids_sincronizados.add(ordem.id)
        lista_numero = existentes_por_numero.setdefault(numero_chave, [])
        if ordem not in lista_numero:
            lista_numero.append(ordem)

    chaves_da_planilha = {item["chave_importacao"] for item in ordens}
    arquivadas = 0
    for ordem in existentes:
        if ordem.origem == "excel" and str(ordem.chave_importacao or "").startswith("excel-ordem:") and ordem.chave_importacao not in chaves_da_planilha and ordem.presente_planilha:
            ordem.presente_planilha = False
            ordem.atualizado_em = agora
            session.add(ordem)
            arquivadas += 1
            _registrar(session, ordem.id, usuario_atual.nome, "Ordem arquivada", f"Não consta mais na planilha {arquivo.filename}.")

    session.commit()
    return {
        "message": "Sincronização concluída.", "ordens_criadas": criadas, "ordens_atualizadas": atualizadas,
        "ordens_sem_alteracao": sem_alteracao, "ordens_arquivadas": arquivadas,
    }
